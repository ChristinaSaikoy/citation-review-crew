import tempfile
import unittest
from pathlib import Path

from docx import Document

from citation_review_crew.tools.docx_reader import extract_cited_passages, read_docx


class DocxReaderTests(unittest.TestCase):
    def _write_doc(self, path: Path) -> None:
        doc = Document()
        doc.add_heading("Introduction", level=1)
        doc.add_paragraph("This paragraph has no citation and should be ignored by extraction.")
        doc.add_paragraph("Prior work reports this effect [1].")
        doc.add_paragraph("Two studies support the same claim [2,3].")
        doc.add_heading("Methods", level=1)
        doc.add_paragraph("The protocol follows an earlier method [4-5].")
        doc.add_heading("References", level=1)
        doc.add_paragraph("[1] Example reference one.")
        doc.add_paragraph("[2] Example reference two.")
        doc.save(path)

    def test_read_docx_returns_nonempty_paragraphs_in_order(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "paper.docx"
            self._write_doc(path)

            text = read_docx(path)

            self.assertIn("Introduction", text)
            self.assertIn("Prior work reports this effect [1].", text)
            self.assertLess(text.index("Introduction"), text.index("Methods"))

    def test_extract_cited_passages_keeps_only_cited_body_paragraphs(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "paper.docx"
            self._write_doc(path)

            result = extract_cited_passages(path)

            self.assertIn("# Cited passages by chapter (3 paragraphs extracted)", result)
            self.assertIn("## Introduction", result)
            self.assertIn("## Methods", result)
            self.assertIn("Prior work reports this effect [1].", result)
            self.assertIn("The protocol follows an earlier method [4-5].", result)
            self.assertNotIn("This paragraph has no citation and should be ignored", result)
            self.assertIn("# Reference list (2 entries)", result)
            self.assertIn("Example reference one.", result)
            self.assertIn("Example reference two.", result)


if __name__ == "__main__":
    unittest.main()
